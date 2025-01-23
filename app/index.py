# index.py

from fastapi import FastAPI, Form, HTTPException
from fastapi.responses import FileResponse
from docx import Document
from docx2pdf import convert
import os
import uuid
from datetime import datetime
from num2words import num2words




app = FastAPI()

# Caminhos de diretório
TEMPLATES_DIR = "./templates/"
OUTPUTS_DIR = "./outputs/"


# Certifique-se de que os diretórios existem
os.makedirs(TEMPLATES_DIR, exist_ok=True)
os.makedirs(OUTPUTS_DIR, exist_ok=True)

def formatar_cpf(cpf: str) -> str:
    """Formata o CPF no padrão 000.000.000-00."""
    cpf = cpf.zfill(11)  # Adiciona zeros à esquerda, se necessário
    return f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}"

def formatar_valor(valor: float) -> str:
    """Formata o valor no padrão brasileiro R$ 1.000,00."""
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def gerar_documento_base(
    modelo_nome: str,
    nome_proprietario: str,
    cpf: str,
    cartorio_number: str,
    matricula: str,
    valor: float,
    corretagem: float,
):
    # Caminho do modelo
    modelo_path = os.path.join(TEMPLATES_DIR, modelo_nome)

    # Verifica se o modelo existe
    if not os.path.exists(modelo_path):
        raise HTTPException(status_code=404, detail=f"Modelo '{modelo_nome}' não encontrado.")

    # Carrega o modelo
    documento = Document(modelo_path)

    # Gera a data atual no formato desejado
    meses = {
        "January": "janeiro", "February": "fevereiro", "March": "março",
        "April": "abril", "May": "maio", "June": "junho",
        "July": "julho", "August": "agosto", "September": "setembro",
        "October": "outubro", "November": "novembro", "December": "dezembro"
    }

    data_atual = datetime.now()
    mes_traduzido = meses[data_atual.strftime("%B")]
    data_extenso = data_atual.strftime(f"%d de {mes_traduzido} de %Y").capitalize()   

    # Converte número para extenso (sem ".00" se for inteiro)
    corretagem_ex = (
        num2words(int(corretagem), lang="pt_BR").capitalize()
        if corretagem.is_integer()
        else num2words(corretagem, lang="pt_BR").capitalize()
    )

    # Formata o CPF e o valor
    cpf_mask = formatar_cpf(cpf)
    valor_mask = formatar_valor(valor)

    # Dicionário de referências a substituir
    refAutorizacao = {
        "{nome_proprietario}": nome_proprietario,
        "{cpf_mask}": cpf_mask,
        "{cartorio_number}": cartorio_number,
        "{matricula}": matricula,
        "{valor_mask}": valor_mask,
        "{corretagem}": f"{corretagem:.0f}" if corretagem.is_integer() else f"{corretagem:.2f}",
        "{corretagem_ex}": corretagem_ex,
        "{data}": data_extenso,
    }

    # Substitui os placeholders no documento
    for paragrafo in documento.paragraphs:
        for codigo, valor in refAutorizacao.items():
            if codigo in paragrafo.text:
                paragrafo.text = paragrafo.text.replace(codigo, valor)

    # Salva o documento preenchido
    unique_id = uuid.uuid4()
    docx_filename = os.path.join(OUTPUTS_DIR, f"{modelo_nome.split('.')[0]}-{unique_id}.docx")
    documento.save(docx_filename)

    # Converte para PDF
    pdf_filename = docx_filename.replace(".docx", ".pdf")
    convert(docx_filename)

    # Retorna o nome do arquivo gerado
    return pdf_filename.split("/")[-1]  # Retorna apenas o nome do arquivo

@app.post("/gerar-documento/autorizacao")
async def gerar_autorizacao(
    nome_proprietario: str = Form(...),
    cpf: str = Form(...),
    cartorio_number: str = Form(...),
    matricula: str = Form(...),
    valor: float = Form(...),
    corretagem: float = Form(...),
):
    arquivo = gerar_documento_base(
        "autorizacao-de-venda.docx",
        nome_proprietario,
        cpf,
        cartorio_number,
        matricula,
        valor,
        corretagem,
    )
    return {
        "msg": "Autorização criada com sucesso!",
        "arquivo_nome": arquivo}

@app.get("/download/{arquivo_nome}")
async def download_arquivo(arquivo_nome: str):
    """
    Permite o download do arquivo gerado ao informar o nome.
    """
    arquivo_path = os.path.join(OUTPUTS_DIR, arquivo_nome)
    if not os.path.exists(arquivo_path):
        raise HTTPException(status_code=404, detail="Arquivo não encontrado.")
    return FileResponse(arquivo_path, media_type="application/pdf", filename=arquivo_nome)

@app.post("/gerar-documento/contrato")
async def gerar_contrato(nome_proprietario: str = Form(...), number_extenso: float = Form(...)):
    return gerar_documento_base("contrato.docx", nome_proprietario, number_extenso)

@app.post("/gerar-documento/termo")
async def gerar_termo(nome_proprietario: str = Form(...), number_extenso: float = Form(...)):
    return gerar_documento_base("termo.docx", nome_proprietario, number_extenso)
